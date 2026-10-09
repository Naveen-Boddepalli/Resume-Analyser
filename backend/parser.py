"""Resume text extraction (PDF, DOCX, TXT). Works on bytes in memory; nothing is written to disk."""

from __future__ import annotations

import io
import re

MAX_BYTES = 5 * 1024 * 1024
MAX_PAGES = 12
_RESUME_WORDS = (
    "education",
    "experience",
    "skills",
    "project",
    "internship",
    "summary",
    "profile",
    "university",
    "college",
    "degree",
    "certification",
    "employment",
)


class ResumeParseError(ValueError):
    """The document could not be read or contained no text."""


def _pdf_text(content: bytes) -> str:
    import fitz  # PyMuPDF

    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:  # noqa: BLE001
        raise ResumeParseError("The PDF could not be opened.") from exc
    parts: list[str] = []
    with doc:
        for page in list(doc)[:MAX_PAGES]:
            blocks = sorted(page.get_text("blocks"), key=lambda b: (round(b[1] / 8), b[0]))
            parts.extend(str(b[4]).strip() for b in blocks if str(b[4]).strip())
    return "\n".join(parts)


def _docx_text(content: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001
        raise ResumeParseError("The DOCX file could not be opened.") from exc
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(c.text for c in row.cells if c.text.strip())
    return "\n".join(parts)


def extract_text(filename: str, content: bytes) -> str:
    """Return the plain text of an uploaded resume."""
    if len(content) > MAX_BYTES:
        raise ResumeParseError(f"File is larger than {MAX_BYTES // (1024 * 1024)} MB.")
    name = (filename or "").lower()
    if name.endswith(".pdf") or content[:5] == b"%PDF-":
        text = _pdf_text(content)
    elif name.endswith(".docx") or content[:2] == b"PK":
        text = _docx_text(content)
    elif name.endswith((".txt", ".md")):
        text = content.decode("utf-8", errors="replace")
    else:
        raise ResumeParseError("Unsupported file type. Upload a PDF, DOCX or TXT file.")
    text = re.sub(r"[ \t]+", " ", text).strip()
    if len(text.split()) < 20:
        raise ResumeParseError("No readable text was found. Scanned/image-only resumes are not supported.")
    return text


def looks_like_resume(text: str) -> bool:
    low = text.lower()
    return sum(1 for w in _RESUME_WORDS if w in low) >= 2
